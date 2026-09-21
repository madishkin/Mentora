from io import BytesIO
from pptx import Presentation

def build_pptx(slides: list[dict], title: str) -> bytes:
    """
    Builds a PPTX presentation from a list of slide dicts.
    Each slide dict should have: 'title', 'points' (list), and optional 'notes'.
    """
    prs = Presentation()
    
    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title_shape = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title_shape.text = title
    subtitle.text = "Сгенерировано в Mentora AI"
    
    # Content slides
    bullet_slide_layout = prs.slide_layouts[1]
    for slide_data in slides:
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        
        # Title
        title_shape = shapes.title
        if title_shape and 'title' in slide_data:
            title_shape.text = slide_data['title']
            
        # Body (bullet points)
        body_shape = shapes.placeholders[1]
        if 'points' in slide_data and slide_data['points']:
            tf = body_shape.text_frame
            for i, pt in enumerate(slide_data['points']):
                if i == 0:
                    p = tf.paragraphs[0]
                    p.text = pt
                else:
                    p = tf.add_paragraph()
                    p.text = pt
                    
        # Notes
        if 'notes' in slide_data and slide_data['notes']:
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = slide_data['notes']
            
    stream = BytesIO()
    prs.save(stream)
    return stream.getvalue()
